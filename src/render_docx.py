from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from extract import SamplePage, iter_field_items


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn(f"w:{key}"), str(value))


def apply_table_grid(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_border(
                cell,
                top={"val": "single", "sz": 8, "space": 0, "color": "666666"},
                bottom={"val": "single", "sz": 8, "space": 0, "color": "666666"},
                left={"val": "single", "sz": 8, "space": 0, "color": "666666"},
                right={"val": "single", "sz": 8, "space": 0, "color": "666666"},
            )


def format_cell(cell, label: str, value: str, min_lines: int = 1) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    label_run = paragraph.add_run(f"{label}\n")
    label_run.bold = True
    label_run.font.size = Pt(8)
    value_run = paragraph.add_run(value if value else "_" * 24)
    value_run.font.size = Pt(9)
    if min_lines > 1:
        for _ in range(min_lines - 1):
            paragraph.add_run("\n")


def set_col_widths(table, widths_cm: list[float]) -> None:
    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)


def add_field_table(document: Document, items: list[tuple[str, str]], widths_cm: list[float], line_counts: list[int] | None = None):
    table = document.add_table(rows=1, cols=len(items))
    apply_table_grid(table)
    set_col_widths(table, widths_cm)
    for index, (label, value) in enumerate(items):
        format_cell(table.rows[0].cells[index], label, value, min_lines=(line_counts[index] if line_counts else 1))
    return table


def add_options(document: Document, title: str, options: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    apply_table_grid(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(18.5)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(f"{title}: ")
    label_run.bold = True
    label_run.font.size = Pt(8)
    value_run = paragraph.add_run(" ".join(options))
    value_run.font.size = Pt(9)


def add_text_area(document: Document, label: str, lines: int = 4) -> None:
    table = document.add_table(rows=1, cols=1)
    apply_table_grid(table)
    cell = table.rows[0].cells[0]
    cell.width = Cm(18.5)
    format_cell(cell, label, "", min_lines=lines)


def add_sample_page(document: Document, sample: SamplePage) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(sample.get("exame_solicitado") or "Mapa de Trabalho GAL")
    run.bold = True
    run.font.size = Pt(11)

    add_field_table(
        document,
        iter_field_items(sample, ["requisicao", "origem", "data_cadastro", "data_primeiros_sintomas", "codigo"]),
        [2.3, 4.0, 3.0, 4.2, 2.1],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["paciente", "idade", "sexo", "idade_gestacional"]),
        [8.0, 2.4, 3.0, 4.2],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["requisitante", "municipio", "profissional_saude"]),
        [6.8, 4.5, 6.3],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["vacinacao", "data_ultima_dose", "qual_vacina"]),
        [6.8, 5.0, 5.8],
    )
    add_field_table(
        document,
        iter_field_items(
            sample,
            ["amostra", "material", "localizacao_amostra", "material_clinico", "data_coleta", "hora_coleta"],
        ),
        [2.2, 3.0, 3.1, 3.2, 3.0, 2.0],
        line_counts=[1, 2, 2, 2, 1, 1],
    )
    add_field_table(
        document,
        iter_field_items(
            sample,
            [
                "data_recebimento",
                "usou_medicamento",
                "qual_medicamento",
                "data_inicio_medicamento",
                "restrito",
                "motivo_restricao",
            ],
        ),
        [2.4, 2.7, 3.1, 3.2, 2.0, 3.1],
        line_counts=[1, 2, 2, 2, 2, 2],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["ri", "kit_lote", "reteste"]),
        [5.0, 9.2, 4.3],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["exame_realizado", "nao_conformidade"]),
        [8.0, 10.5],
        line_counts=[2, 2],
    )
    add_field_table(
        document,
        iter_field_items(sample, ["exame_nao_realizado", "exame_cancelado"]),
        [9.25, 9.25],
    )
    add_options(document, "DNA para Mycobacterium tuberculosis", sample.option_groups["dna_mt"])
    add_options(document, "Rifampicina", sample.option_groups["rifampicina"])
    add_options(document, "Complemento", sample.option_groups["complemento"])
    add_options(document, "Aspecto da Amostra de Escarro", sample.option_groups["aspecto_amostra"])
    add_text_area(document, "Valor de Referência", lines=4)
    add_text_area(document, "Observações", lines=5)


def build_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)
    section.top_margin = Cm(1.0)
    section.bottom_margin = Cm(1.0)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(9)
    return document


def render_docx(samples: list[SamplePage], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    document = build_document()

    for index, sample in enumerate(samples):
        if index > 0:
            document.add_page_break()
        add_sample_page(document, sample)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return output_path
